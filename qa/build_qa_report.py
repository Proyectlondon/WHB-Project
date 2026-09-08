from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / "docs" / "auditoria_web_whb_project_2026_09_07.docx"

GREEN = "123A2B"
GOLD = "D69C2D"
PALE = "F3F5F2"
GRID = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tbl_pr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), GRID)
        node.append(item)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths):
    table.autofit = False
    borders(table)
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)
            if row_index == 0:
                shade(cell, GREEN)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                    run.font.size = Pt(9)
            elif row_index % 2 == 0:
                shade(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    if row_index > 0:
                        run.font.size = Pt(8.2)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)


def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True
    return heading


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].font.color.rgb = RGBColor(32, 38, 34)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Aptos Display"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(28)
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
styles["Heading 1"].font.size = Pt(17)
styles["Heading 1"].paragraph_format.space_before = Pt(15)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].font.size = Pt(12)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

title = doc.add_paragraph(style="Title")
title.add_run("Auditoría web de WHB Project")
subtitle = doc.add_paragraph()
subtitle.add_run("Informe de calidad y verificación de producción").bold = True
subtitle.add_run("\n7 de septiembre de 2026 · America Bogotá")

opening = doc.add_paragraph()
opening.add_run("Resultado final. ").bold = True
opening.add_run(
    "La versión pública en https://whb-project.vercel.app superó la batería automatizada después de corregir la publicación incompleta de los archivos de audio. "
    "Las 46 pistas están íntegras, la introducción reproduce, las interacciones móviles responden y no quedan errores de consola ni incumplimientos serios de accesibilidad WCAG A o AA en el recorrido probado."
)

add_heading(doc, "Alcance y fuentes", 1)
doc.add_paragraph(
    "La revisión cubrió la experiencia pública y la copia local de site-v1.5. Se comprobó el catálogo de música, el video de introducción, la navegación por secciones, el carrusel editorial, el gesto táctil, el minirreproductor, el formulario de contacto, los recursos visuales y los nombres accesibles de los controles."
)
for item in (
    "Sitio público: https://whb-project.vercel.app",
    "Repositorio: https://github.com/Proyectlondon/WHB-Project",
    "Fuente editorial y multimedia: site-v1.5/content/catalog.json",
    "Pruebas reproducibles: qa/media-integrity.cjs, qa/whb-e2e.cjs y qa/full-site-audit.cjs",
):
    add_bullet(doc, item)

add_heading(doc, "Causa raíz y corrección", 1)
doc.add_paragraph(
    "Cada archivo de audio publicado mediante la interfaz anterior quedó recortado a unos 786 KB. El encabezado MP3 conservó la duración original, pero el contenido físico terminaba después de cerca de 16 segundos. Por eso el navegador saltaba de ese punto al final declarado de la canción. La falla afectaba las 46 pistas y también tres videos grandes."
)
doc.add_paragraph(
    "Se reemplazaron los binarios dañados con las copias locales completas mediante Git nativo. Los audios suman 264,209,169 bytes y contienen entre 2:42 y 5:30 de material reproducible. El sitio ahora solicita los medios desde el mismo dominio de Vercel. El service worker excluye audio, video y respuestas parciales de la caché para evitar que una petición Range vuelva a almacenarse como archivo completo."
)

add_heading(doc, "Resultados de las pruebas", 1)
rows = [
    ("Integridad local", "46 pistas", "Aprobado", "Todos los MP3 superan 1 MB, 1,000 tramas y 60 segundos reales"),
    ("Integridad pública", "64 recursos", "Aprobado", "46 audios, 17 imágenes y 1 video responden sin faltantes"),
    ("Reproducción", "Caminos de Zipacón", "Aprobado", "Avanza de 0:00 a 0:22 en escritorio y móvil; duración 3:55"),
    ("Introducción", "390 por 844", "Aprobado", "El video inicia y el tiempo avanza antes de entrar al sitio"),
    ("Interacción móvil", "Carrusel editorial", "Aprobado", "Botones y gesto lateral cambian de canción"),
    ("Catálogo", "46 audios y 22 videos", "Aprobado", "Conteos completos e IDs de YouTube únicos"),
    ("Accesibilidad", "WCAG A y AA", "Aprobado", "Sin hallazgos serios o críticos en Axe"),
    ("Estabilidad", "Consola", "Aprobado", "Sin errores de JavaScript en los recorridos finales"),
]
table = doc.add_table(rows=1, cols=4)
for idx, value in enumerate(("Prueba", "Cobertura", "Estado", "Evidencia")):
    table.rows[0].cells[idx].text = value
for record in rows:
    cells = table.add_row().cells
    for idx, value in enumerate(record):
        cells[idx].text = value
format_table(table, [Inches(1.32), Inches(1.42), Inches(0.88), Inches(3.2)])

add_heading(doc, "Hallazgos y resolución", 1)
findings = [
    ("QA 01", "Crítico", "Las canciones terminaban después de unos 16 segundos", "Se restauraron las 46 copias completas", "Cerrado"),
    ("QA 02", "Alto", "Los binarios grandes se recortaban durante la publicación", "La entrega se realizó con Git nativo y se comprobó el tamaño remoto", "Cerrado"),
    ("QA 03", "Alto", "La caché podía almacenar respuestas parciales de medios", "El service worker omite medios y peticiones Range; caché v2", "Cerrado"),
    ("QA 04", "Medio", "El minirreproductor podía conservar el título anterior", "El estado se sincroniza al seleccionar cada pista", "Cerrado"),
    ("QA 05", "Medio", "La etiqueta Integrante histórico no alcanzaba contraste AA", "Se ajustó el color sin perder la jerarquía visual", "Cerrado"),
    ("QA 06", "Medio", "La prueba táctil necesitaba validar el gesto real", "Se añadió una simulación de PointerEvent táctil", "Cerrado"),
    ("QA 07", "Bajo", "Faltaba una prueba automatizada repetible de medios", "Se añadieron tres scripts de regresión al repositorio", "Cerrado"),
]
table = doc.add_table(rows=1, cols=5)
for idx, value in enumerate(("ID", "Impacto", "Hallazgo", "Resolución", "Estado")):
    table.rows[0].cells[idx].text = value
for record in findings:
    cells = table.add_row().cells
    for idx, value in enumerate(record):
        cells[idx].text = value
format_table(table, [Inches(0.55), Inches(0.72), Inches(2.0), Inches(2.4), Inches(0.72)])

add_heading(doc, "Cinco controles para futuras versiones", 1)
controls = [
    "Ejecutar qa/media-integrity.cjs antes de cada publicación que incluya audios.",
    "Ejecutar qa/whb-e2e.cjs contra la vista previa y contra el dominio público.",
    "Ejecutar qa/full-site-audit.cjs después de que Vercel marque el despliegue como READY.",
    "Evitar publicar audio o video grande mediante herramientas que serialicen el binario dentro de una solicitud de texto.",
    "Mantener audio y video fuera de la caché del service worker salvo que se implemente manejo explícito de peticiones Range.",
]
for item in controls:
    add_bullet(doc, item)

add_heading(doc, "Matriz de aceptación", 1)
acceptance = [
    ("Escritorio", "1440 por 900", "Reproducción, scroll, minirreproductor, consola", "Aprobado"),
    ("Móvil", "390 por 844", "Reproducción, gesto lateral, navegación y scroll", "Aprobado"),
    ("Contenido", "Catálogo completo", "46 audios, 22 videos y 17 imágenes", "Aprobado"),
    ("Accesibilidad", "WCAG A y AA", "Contraste y nombres accesibles", "Aprobado"),
    ("Producción", "Vercel", "Despliegue cdae2d3 y dominio principal", "READY"),
]
table = doc.add_table(rows=1, cols=4)
for idx, value in enumerate(("Superficie", "Configuración", "Criterio", "Resultado")):
    table.rows[0].cells[idx].text = value
for record in acceptance:
    cells = table.add_row().cells
    for idx, value in enumerate(record):
        cells[idx].text = value
format_table(table, [Inches(1.0), Inches(1.35), Inches(3.45), Inches(0.85)])

add_heading(doc, "Estado de entrega", 1)
doc.add_paragraph(
    "GitHub contiene los archivos completos en el commit cdae2d38e83770bd93348822ebfa79d529a483ef. Vercel publicó ese commit en producción con estado READY. La auditoría final del dominio principal terminó sin hallazgos, errores de consola, medios incompletos ni violaciones serias de accesibilidad en el recorrido probado. Estas pruebas deben complementarse con revisiones periódicas en dispositivos físicos y redes móviles lentas."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
